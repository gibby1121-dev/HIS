"""Mid-Iowa Auction's own data: the edge an operator doesn't have.

Two exports MIA already pulls (see the post-sale-auction-recap and
tractorhouse-vault skills):

* **Auction Summary Report** (TractorHouse/HiBid post-sale .xlsx). The
  ``Details`` sheet, data from row 6, 0-based columns: 0 ListingID, 1 Lot,
  3 Year, 4 Manufacturer, 5 Model, 9 Auction Value (Sandhills estimate),
  10 Market Value, 12 Final price *with* buyer's premium, 14 Hammer.
* **ExportFleet** (Sandhills dealer account): ``VIP+ AuctionValue``,
  ``VIP+ WholesaleValue``, ``VIP+ MarketValue``, ``VIP+ AskingValue`` per
  serial number.

``calibrate`` measures how MIA hammer prices actually land against the
Sandhills auction estimate, per category, so a VIP+ auction value for an
operator's trade unit can be turned into an *MIA-realized* expectation with
an empirical range — instead of trusting a national model blindly.
"""

from __future__ import annotations

import glob
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd

from .categories import OTHER, classify

# Column indices in the Auction Summary Report "Details" sheet.
_COL = {"listing": 0, "lot": 1, "year": 3, "make": 4, "model": 5,
        "auction_value": 9, "market_value": 10, "final_with_bp": 12, "hammer": 14}
_DATA_START_ROW = 5  # 0-based; data starts on spreadsheet row 6

# Shrink small-sample ratios toward 1.0 (trust Sandhills) with this many
# pseudo-observations, so two lucky lots can't swing a category.
SHRINK_K = 5
MIN_LOTS_FOR_CATEGORY = 3


class MiaDataError(ValueError):
    pass


def _num(x):
    try:
        return float(str(x).replace(",", "").replace("$", "").strip())
    except (TypeError, ValueError):
        return None


def load_auction_summary(path: Path, sale_date: str | None = None) -> pd.DataFrame:
    """Parse one Auction Summary Report into one row per lot."""
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover
        raise MiaDataError("openpyxl is required: pip install openpyxl") from exc
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise MiaDataError(f"Could not open '{path}': {exc}") from exc
    if "Details" not in wb.sheetnames:
        raise MiaDataError(f"'{path}' has no 'Details' sheet (found {wb.sheetnames}).")
    rows = list(wb["Details"].iter_rows(values_only=True))[_DATA_START_ROW:]
    out = []
    for r in rows:
        if not r or len(r) <= _COL["hammer"] or not r[_COL["listing"]]:
            continue
        make = str(r[_COL["make"]] or "").strip()
        model = str(r[_COL["model"]] or "").strip()
        out.append({
            "Source": Path(path).name,
            "SaleDate": sale_date,
            "Lot": r[_COL["lot"]],
            "Year": _num(r[_COL["year"]]),
            "Make": make,
            "Model": model,
            "Category": classify(make, model),
            "AuctionValue": _num(r[_COL["auction_value"]]),
            "MarketValue": _num(r[_COL["market_value"]]),
            "FinalWithBP": _num(r[_COL["final_with_bp"]]),
            "Hammer": _num(r[_COL["hammer"]]),
        })
    if not out:
        raise MiaDataError(f"'{path}' Details sheet has no lot rows.")
    return pd.DataFrame(out)


def load_auction_history(folder: Path) -> pd.DataFrame:
    """Load every ``Auction Summary Report*.xlsx`` in a folder."""
    files = sorted(glob.glob(str(Path(folder) / "*Auction Summary*.xlsx")))
    if not files:
        raise MiaDataError(f"No 'Auction Summary Report*.xlsx' files in {folder}.")
    return pd.concat([load_auction_summary(Path(f)) for f in files], ignore_index=True)


@dataclass
class CategoryCalibration:
    category: str
    lots: int
    median_ratio: float      # hammer / Sandhills auction estimate (raw)
    p25_ratio: float
    p75_ratio: float
    p10_ratio: float
    p90_ratio: float
    applied_ratio: float     # shrunk toward 1.0 by sample size
    implied_bp_pct: float | None  # observed (final-with-BP / hammer) - 1


@dataclass
class Calibration:
    by_category: dict[str, CategoryCalibration]
    overall: CategoryCalibration | None
    lots_used: int
    lots_skipped: int
    # category -> (raw median % error, calibrated median % error), held-out
    validation: dict[str, tuple[float, float]] = field(default_factory=dict)

    def for_category(self, category: str) -> CategoryCalibration | None:
        """Calibration for a category, only if it beat raw Sandhills on held-out lots.

        No pooled fallback: a backtest showed a pooled ratio makes a category
        with its own bias (sprayers ran below estimate) worse, not better.
        """
        c = self.by_category.get(category)
        if not c or c.lots < MIN_LOTS_FOR_CATEGORY:
            return None
        v = self.validation.get(category)
        if v is None or v[1] >= v[0]:
            return None
        return c


def _summarize(cat: str, df: pd.DataFrame) -> CategoryCalibration:
    r = df["Hammer"] / df["AuctionValue"]
    n = len(r)
    med = float(r.median())
    bp = (df["FinalWithBP"] / df["Hammer"] - 1).dropna()
    bp = bp[(bp >= 0) & (bp < 0.25)]
    return CategoryCalibration(
        category=cat,
        lots=n,
        median_ratio=med,
        p25_ratio=float(r.quantile(0.25)),
        p75_ratio=float(r.quantile(0.75)),
        p10_ratio=float(r.quantile(0.10)),
        p90_ratio=float(r.quantile(0.90)),
        applied_ratio=(n * med + SHRINK_K * 1.0) / (n + SHRINK_K),
        implied_bp_pct=float(bp.median() * 100) if len(bp) else None,
    )


def calibrate(lots: pd.DataFrame, min_estimate: float = 25_000) -> Calibration:
    """Hammer-vs-Sandhills-estimate ratios from MIA's realized sales.

    Only lots with both a Sandhills auction estimate and a hammer price are
    used, and only above ``min_estimate`` — small-ticket lots behave
    differently and this product is about six-figure iron.
    """
    usable = lots[
        lots["AuctionValue"].fillna(0).ge(min_estimate) & lots["Hammer"].fillna(0).gt(0)
    ]
    by_cat = {
        cat: _summarize(cat, g)
        for cat, g in usable.groupby("Category")
        if cat != OTHER
    }
    overall = _summarize("All large iron", usable) if len(usable) else None
    return Calibration(by_cat, overall, len(usable), len(lots) - len(usable),
                       _validate(usable))


def _validate(usable: pd.DataFrame) -> dict[str, tuple[float, float]]:
    """Leave-one-out per category: raw vs calibrated median absolute % error."""
    out = {}
    for cat, g in usable.groupby("Category"):
        if cat == OTHER or len(g) < MIN_LOTS_FOR_CATEGORY + 1:
            continue
        raw, cal = [], []
        for idx in g.index:
            rest = g.drop(index=idx)
            ratio = _summarize(cat, rest).applied_ratio
            r = g.loc[idx]
            raw.append(abs(r["AuctionValue"] / r["Hammer"] - 1) * 100)
            cal.append(abs(r["AuctionValue"] * ratio / r["Hammer"] - 1) * 100)
        out[cat] = (float(pd.Series(raw).median()), float(pd.Series(cal).median()))
    return out


def load_vip_values(export_fleet: Path, serial: str) -> dict | None:
    """VIP+ values for one serial from a Sandhills ExportFleet csv/xlsx."""
    p = Path(export_fleet)
    if p.suffix.lower() in (".xlsx", ".xls"):
        df = pd.read_excel(p, dtype=str)
    else:
        df = pd.read_csv(p, dtype=str, encoding="utf-8-sig")
    if "VINSerialNumber" not in df.columns:
        raise MiaDataError(f"'{p.name}' has no VINSerialNumber column.")
    hit = df[df["VINSerialNumber"].str.strip().str.upper() == serial.strip().upper()]
    if hit.empty:
        return None
    r = hit.iloc[0]
    pick = lambda c: _num(r.get(c)) if c in r else None  # noqa: E731
    return {
        "auction": pick("VIP+ AuctionValue"),
        "wholesale": pick("VIP+ WholesaleValue"),
        "market": pick("VIP+ MarketValue"),
        "asking": pick("VIP+ AskingValue"),
    }
